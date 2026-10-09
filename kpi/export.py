from io import BytesIO

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kpi.services import GRADE_LABELS, OUTCOME_LABELS
from tasks.models import Task

HEADER_FILL = PatternFill("solid", fgColor="DCE7F5")


def _write_table(ws, start_row, headers, rows):
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, values in enumerate(rows, start_row + 1):
        for c, value in enumerate(values, 1):
            ws.cell(row=r, column=c, value=value)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def build_workbook(month_label, rows, period, organization=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws["A1"] = (organization + " — " if organization else "") + _("Employee KPI for %(month)s") % {
        "month": month_label
    }
    ws["A1"].font = Font(bold=True, size=13)
    status = (
        _("Finalised on %(date)s") % {"date": timezone.localtime(period.finalized_at).strftime("%d.%m.%Y")}
        if period
        else _("Preliminary (month not finalised)")
    )
    ws["A2"] = status + " · " + _("Generated %(date)s") % {
        "date": timezone.localtime().strftime("%d.%m.%Y %H:%M")
    }
    headers = [
        "№",
        _("Employee"),
        _("Position"),
        _("Department"),
        _("Tasks assessed"),
        _("Completed"),
        _("On time"),
        _("Not done by deadline"),
        _("Awaiting review"),
        _("Execution, %"),
        _("Timeliness, %"),
        _("Quality, %"),
        _("Workload, %"),
        _("Base score"),
        _("Adjustment"),
        _("KPI"),
        _("Grade"),
    ]
    data = [
        [
            i,
            row["name"],
            row["position"],
            row["department"],
            row["assessed"],
            row["done"],
            row["on_time"],
            row["failed"],
            row["pending"],
            row["completion"],
            row["timeliness"],
            row["quality"],
            row["workload"],
            row["base_score"],
            row["adjustment"],
            row["score"],
            str(GRADE_LABELS[row["grade"]]) if row["grade"] else "",
        ]
        for i, row in enumerate(rows, 1)
    ]
    _write_table(ws, 4, headers, data)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 26

    detail = wb.create_sheet(_("Tasks"))
    complexity = dict(Task.Complexity.choices)
    headers = [
        _("Employee"),
        "ID",
        _("Task"),
        _("Complexity"),
        _("Points"),
        _("Deadline"),
        _("Submitted"),
        _("Result"),
        _("Quality rating"),
    ]
    data = [
        [
            row["name"],
            task["id"],
            task["title"],
            str(complexity[task["complexity"]]),
            task["points"],
            task["deadline"],
            task["submitted"] or "",
            str(OUTCOME_LABELS[task["outcome"]]),
            task["rating"] or "",
        ]
        for row in rows
        for task in row["tasks"]
    ]
    _write_table(detail, 1, headers, data)
    detail.column_dimensions["A"].width = 30
    detail.column_dimensions["C"].width = 60

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
